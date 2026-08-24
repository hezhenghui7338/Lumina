"""Import + summarize spike tests."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from lumina_core.config import (
    CHUNKER_VERSION,
    RESEGMENT_MAX_TARGET_CHARS,
    RESEGMENT_MIN_TARGET_CHARS,
    Settings,
)
from lumina_core.db.repos import BookRepo
from lumina_core.main import create_app
from lumina_core.models.router import set_router
from tests.support.import_helpers import import_sample_book, wait_for_ingest
from tests.support.mock_router import MockModelRouter, load_json_fixture

pytestmark = pytest.mark.e2e

LLM_FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "llm"
BOOK_FIXTURES = Path(__file__).resolve().parents[1] / "fixtures" / "books"


def _write_extended_format_fixture(tmp_path: Path, extension: str) -> Path:
    path = tmp_path / f"extended-{extension}.{extension}"
    title = f"{extension.upper()} 测试书"
    body = f"{extension.upper()} 格式正文段落。"
    if extension == "md":
        path.write_text(f"# {title}\n\n{body}", encoding="utf-8")
    elif extension == "html":
        path.write_text(
            f"<html><head><title>{title}</title></head>"
            f"<body><h1>第一章</h1><p>{body}</p></body></html>",
            encoding="utf-8",
        )
    elif extension == "rtf":
        path.write_text(
            rf"{{\rtf1\ansi{{\info{{\title {title}}}}}\b Chapter\b0\par {body}}}",
            encoding="utf-8",
        )
    elif extension == "docx":
        from docx import Document

        document = Document()
        document.core_properties.title = title
        document.add_heading("第一章", level=1)
        document.add_paragraph(body)
        document.save(path)
    elif extension == "odt":
        from odf import dc, text
        from odf.opendocument import OpenDocumentText

        document = OpenDocumentText()
        document.meta.addElement(dc.Title(text=title))
        document.text.addElement(text.H(outlinelevel=1, text="第一章"))
        document.text.addElement(text.P(text=body))
        document.save(str(path))
    elif extension == "fb2":
        path.write_text(
            f"""<?xml version="1.0" encoding="utf-8"?>
            <FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0">
              <description><title-info><book-title>{title}</book-title></title-info></description>
              <body><section><title><p>第一章</p></title><p>{body}</p></section></body>
            </FictionBook>""",
            encoding="utf-8",
        )
    else:
        raise AssertionError(f"Unhandled fixture extension: {extension}")
    return path


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("LUMINA_DATA_DIR", str(tmp_path))
    router = MockModelRouter(
        responses={
            "summarize": load_json_fixture(LLM_FIXTURES / "summary_segment0.json"),
            "chat": load_json_fixture(LLM_FIXTURES / "chat_with_citation.json"),
            "translate": "示例译文。",
        }
    )
    app = create_app(Settings(data_dir=tmp_path))
    app.state.lumina.router = router
    app.state.lumina.job_queue.router = router
    set_router(router)
    with TestClient(app) as c:
        yield c


def test_health(client):
    health = client.get("/health").json()
    assert health["status"] == "ok"
    assert health["pid"] > 1
    assert health["chunker_version"] == CHUNKER_VERSION


def test_import_returns_processing_immediately(client):
    sample = BOOK_FIXTURES / "sample.txt"
    resp = client.post("/books/import", json={"paths": [str(sample)]})
    assert resp.status_code == 200
    body = resp.json()["books"][0]
    assert body["status"] == "processing"
    book_id = body["book_id"]
    listed = client.get("/books").json()["books"]
    assert any(b["id"] == book_id for b in listed)
    finished = wait_for_ingest(client, book_id)
    assert finished["status"] in ("unread", "reading", "summarized")
    assert finished.get("segment_count", 0) > 0


@pytest.mark.parametrize("extension", ["md", "html", "rtf", "docx", "odt", "fb2"])
def test_import_extended_text_formats(client, tmp_path, extension):
    sample = _write_extended_format_fixture(tmp_path, extension)
    response = client.post("/books/import", json={"paths": [str(sample)]})
    assert response.status_code == 200
    imported = response.json()["books"][0]
    assert imported["status"] == "processing"

    finished = wait_for_ingest(client, imported["book_id"])
    assert finished["status"] in ("unread", "reading", "summarized")
    assert finished["segment_count"] > 0
    first_segment = client.get(f"/books/{imported['book_id']}/segments/0").json()
    assert extension.upper() in first_segment["raw_text"]


def test_import_txt_triggers_prefetch(client, tmp_path):
    assert client.put("/settings", json={"auto_start_summary": True}).status_code == 200

    sample = BOOK_FIXTURES / "sample.txt"
    resp = client.post("/books/import", json={"paths": [str(sample)]})
    assert resp.status_code == 200
    book_id = resp.json()["books"][0]["book_id"]
    wait_for_ingest(client, book_id)

    import time

    for _ in range(50):
        segs = client.get(f"/books/{book_id}/segments").json()["segments"]
        if segs and segs[0]["summary_status"] == "ready":
            break
        time.sleep(0.1)

    # List API is meta-only (never-freeze): no raw_text payload.
    list_seg = client.get(f"/books/{book_id}/segments").json()["segments"][0]
    assert "raw_text" not in list_seg or list_seg.get("raw_text") is None

    seg0 = client.get(f"/books/{book_id}/segments/0").json()
    assert seg0["summary_status"] == "ready"
    assert seg0.get("raw_text")
    summary = json.loads(seg0["summary_json"])
    assert "sentences" in summary
    assert seg0["label"]


def test_import_does_not_prefetch_when_auto_start_off(client, tmp_path):
    assert client.get("/settings").json()["auto_start_summary"] is False

    sample = BOOK_FIXTURES / "sample.txt"
    book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0]["book_id"]
    wait_for_ingest(client, book_id)

    import time

    time.sleep(0.3)
    segs = client.get(f"/books/{book_id}/segments").json()["segments"]
    assert segs
    assert all(s["summary_status"] == "pending" for s in segs)


def test_import_duplicate_returns_409(client, tmp_path):
    sample = BOOK_FIXTURES / "sample.txt"
    book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0]["book_id"]
    wait_for_ingest(client, book_id)
    resp = client.post("/books/import", json={"paths": [str(sample)]})
    assert resp.status_code == 409
    assert "existing_book_id" in resp.json()["detail"]


def test_import_overwrite_purges_old_data(client, tmp_path):
    sample = BOOK_FIXTURES / "sample.txt"
    old_book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0]["book_id"]
    wait_for_ingest(client, old_book_id)

    segments = client.get(f"/books/{old_book_id}/segments").json()["segments"]
    assert segments
    seg_id = segments[0]["id"]
    note_resp = client.post(
        "/notes",
        json={
            "book_id": old_book_id,
            "segment_id": seg_id,
            "content": "overwrite test note",
        },
    )
    assert note_resp.status_code == 200

    old_books_dir = tmp_path / "books" / old_book_id
    assert old_books_dir.exists()

    resp = client.post(
        "/books/import",
        json={"paths": [str(sample)], "overwrite": True},
    )
    assert resp.status_code == 200
    new_book_id = resp.json()["books"][0]["book_id"]
    assert new_book_id != old_book_id

    assert client.get(f"/books/{old_book_id}").status_code == 404
    assert not old_books_dir.exists()

    conn = client.app.state.lumina.conn
    old_seg_count = conn.execute(
        "SELECT COUNT(*) AS c FROM segments WHERE book_id = ?",
        (old_book_id,),
    ).fetchone()["c"]
    old_note_count = conn.execute(
        "SELECT COUNT(*) AS c FROM notes WHERE book_id = ?",
        (old_book_id,),
    ).fetchone()["c"]
    assert old_seg_count == 0
    assert old_note_count == 0

    finished = wait_for_ingest(client, new_book_id)
    assert finished.get("segment_count", 0) > 0
    notes = client.get("/notes", params={"book_id": new_book_id}).json()["notes"]
    assert notes == []


def test_resegment_book_uses_new_size_and_clears_segment_bound_data(client, tmp_path):
    sample = tmp_path / "long-book.txt"
    sample.write_text(
        "\n\n".join(
            f"第 {i} 段。这是一段用于验证整书重新分段的测试文字，包含完整句子和稳定边界。"
            for i in range(700)
        ),
        encoding="utf-8",
    )
    book_id = client.post(
        "/books/import", json={"paths": [str(sample)]}
    ).json()["books"][0]["book_id"]
    imported = wait_for_ingest(client, book_id)
    assert imported["status"] in ("unread", "reading", "summarized"), imported
    assert imported.get("segment_count", 0) > 0, imported
    original_count = imported["segment_count"]

    segments = client.get(f"/books/{book_id}/segments").json()["segments"]
    assert segments, imported
    first_segment = segments[0]
    note = client.post(
        "/notes",
        json={
            "book_id": book_id,
            "segment_id": first_segment["id"],
            "content": "重新分段前的笔记",
        },
    )
    assert note.status_code == 200

    response = client.post(
        f"/books/{book_id}/resegment",
        json={"chunk_target_chars": 1500},
    )
    assert response.status_code == 202
    finished = wait_for_ingest(client, book_id)

    assert finished["segment_count"] > original_count
    assert finished["current_segment_index"] == 0
    assert finished["chunk_target_chars"] == 1500
    assert client.get("/notes", params={"book_id": book_id}).json()["notes"] == []
    segments = client.get(f"/books/{book_id}/segments").json()["segments"]
    assert segments
    assert all(segment["summary_status"] == "pending" for segment in segments)


@pytest.mark.parametrize(
    "target",
    [RESEGMENT_MIN_TARGET_CHARS - 1, RESEGMENT_MAX_TARGET_CHARS + 1],
)
def test_resegment_rejects_out_of_range_size(client, target):
    response = client.post(
        "/books/missing/resegment",
        json={"chunk_target_chars": target},
    )
    assert response.status_code == 422


@pytest.mark.parametrize(
    "target",
    [RESEGMENT_MIN_TARGET_CHARS, RESEGMENT_MAX_TARGET_CHARS],
)
def test_resegment_accepts_boundary_target_size(client, target):
    response = client.post(
        "/books/missing/resegment",
        json={"chunk_target_chars": target},
    )
    assert response.status_code == 404


def test_resegment_rejects_missing_and_processing_books(client):
    missing = client.post(
        "/books/missing/resegment",
        json={"chunk_target_chars": 3000},
    )
    assert missing.status_code == 404

    book_id = import_sample_book(client)
    BookRepo(client.app.state.lumina.conn).update(book_id, status="processing")
    conflict = client.post(
        f"/books/{book_id}/resegment",
        json={"chunk_target_chars": 3000},
    )
    assert conflict.status_code == 409


def test_cancel_resegment_preserves_existing_data(client, monkeypatch):
    import time

    from lumina_core.jobs import resegment as resegment_module

    book_id = import_sample_book(client)
    old_segments = client.get(f"/books/{book_id}/segments").json()["segments"]
    old_segment_ids = [segment["id"] for segment in old_segments]
    note = client.post(
        "/notes",
        json={
            "book_id": book_id,
            "segment_id": old_segment_ids[0],
            "content": "取消后应保留",
        },
    )
    assert note.status_code == 200

    original_load_document = resegment_module.load_document

    def slow_load_document(*args, **kwargs):
        time.sleep(0.3)
        return original_load_document(*args, **kwargs)

    monkeypatch.setattr(resegment_module, "load_document", slow_load_document)
    started = client.post(
        f"/books/{book_id}/resegment",
        json={"chunk_target_chars": 1500},
    )
    assert started.status_code == 202
    assert started.json()["processing_kind"] == "resegment"
    assert client.get(f"/books/{book_id}").json()["processing_kind"] == "resegment"
    cancelled = client.post(f"/books/{book_id}/resegment/cancel", json={})
    assert cancelled.status_code == 202
    assert cancelled.json()["status"] == "cancelling"

    restored = wait_for_ingest(client, book_id)
    assert restored["status"] != "processing"
    current_segments = client.get(f"/books/{book_id}/segments").json()["segments"]
    assert [segment["id"] for segment in current_segments] == old_segment_ids
    notes = client.get("/notes", params={"book_id": book_id}).json()["notes"]
    assert [item["content"] for item in notes] == ["取消后应保留"]


def test_cancel_ingest_marks_book_error(client, monkeypatch):
    import threading
    import time

    from lumina_core.ingest.progress import DocumentLoadCancelled
    from lumina_core.jobs import ingest as ingest_module

    started = threading.Event()

    def slow_load_document(*args, **kwargs):
        started.set()
        cancel = kwargs.get("cancel_event")
        for _ in range(80):
            if cancel is not None and cancel.is_set():
                raise DocumentLoadCancelled("已取消")
            time.sleep(0.05)
        return "正文", {}

    monkeypatch.setattr(ingest_module, "load_document", slow_load_document)
    sample = BOOK_FIXTURES / "sample.txt"
    book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0][
        "book_id"
    ]
    assert started.wait(timeout=2)
    assert client.get(f"/books/{book_id}").json()["processing_kind"] == "ingest"
    cancelled = client.post(f"/books/{book_id}/ingest/cancel", json={})
    assert cancelled.status_code == 202
    assert cancelled.json()["status"] == "cancelling"
    finished = wait_for_ingest(client, book_id, timeout=5.0)
    assert finished["status"] == "error"
    assert finished.get("ingest_error") == "已取消导入"


def test_chat_json_mode(client, tmp_path):
    assert client.put("/settings", json={"auto_start_summary": True}).status_code == 200
    sample = BOOK_FIXTURES / "sample.txt"
    book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0]["book_id"]
    wait_for_ingest(client, book_id)

    import time

    for _ in range(50):
        segs = client.get(f"/books/{book_id}/segments").json()["segments"]
        if segs and segs[0]["summary_status"] == "ready":
            break
        time.sleep(0.1)

    chat = client.post(
        f"/books/{book_id}/chat",
        json={"message": "这段讲了什么？", "segment_index": 0},
    )
    assert chat.status_code == 200
    body = chat.json()
    assert "answer" in body
    assert body["citations"]


def test_book_scope_chat_after_index(client, tmp_path):
    assert client.put("/settings", json={"auto_start_summary": True}).status_code == 200
    sample = BOOK_FIXTURES / "sample.txt"
    book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0]["book_id"]
    wait_for_ingest(client, book_id)

    import time

    # Nothing builds the index in the background, so book-scope chat must refuse.
    too_early = client.post(
        f"/books/{book_id}/chat",
        json={"message": "全书主旨？", "segment_index": 0, "scope": "book"},
    )
    assert too_early.status_code == 409
    assert client.get(f"/books/{book_id}").json()["index_status"] != "ready"

    for _ in range(80):
        segs = client.get(f"/books/{book_id}/segments").json()["segments"]
        if segs and all(s["summary_status"] == "ready" for s in segs):
            break
        time.sleep(0.1)

    assert client.post(f"/books/{book_id}/index").status_code == 200

    book = None
    for _ in range(80):
        book = client.get(f"/books/{book_id}").json()
        if book.get("index_status") == "ready":
            break
        time.sleep(0.1)
    assert book is not None
    assert book["index_status"] == "ready"

    chat = client.post(
        f"/books/{book_id}/chat",
        json={"message": "全书主旨是什么？", "segment_index": 0, "scope": "book"},
    )
    assert chat.status_code == 200
    body = chat.json()
    assert "answer" in body
    assert body["citations"]


def test_export_markdown(client, tmp_path):
    assert client.put("/settings", json={"auto_start_summary": True}).status_code == 200
    sample = BOOK_FIXTURES / "sample.txt"
    book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0]["book_id"]
    wait_for_ingest(client, book_id)
    import time

    for _ in range(50):
        segs = client.get(f"/books/{book_id}/segments").json()["segments"]
        if segs and segs[0]["summary_status"] == "ready":
            break
        time.sleep(0.1)

    resp = client.post(f"/books/{book_id}/export", json={"include_notes": False})
    assert resp.status_code == 200
    assert "摘要版" in resp.text


def test_export_chinese_title_returns_200(client, tmp_path):
    assert client.put("/settings", json={"auto_start_summary": True}).status_code == 200
    sample = BOOK_FIXTURES / "sample.txt"
    book_id = client.post("/books/import", json={"paths": [str(sample)]}).json()["books"][0]["book_id"]
    wait_for_ingest(client, book_id)
    import time

    for _ in range(50):
        segs = client.get(f"/books/{book_id}/segments").json()["segments"]
        if segs and segs[0]["summary_status"] == "ready":
            break
        time.sleep(0.1)

    patch = client.patch(f"/books/{book_id}", json={"title": "三体"})
    assert patch.status_code == 200

    resp = client.post(f"/books/{book_id}/export", json={"include_notes": False})
    assert resp.status_code == 200
    assert "摘要版" in resp.text
    disposition = resp.headers.get("Content-Disposition", "")
    assert "filename" in disposition
    disposition.encode("latin-1")


def test_settings_roundtrip(client):
    resp = client.put(
        "/settings",
        json={
            "target_language": "en-US",
            "web_search_provider": "tavily",
            "tavily_api_key": "tvly-secret",
        },
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["target_language"] == "en-US"
    assert body["web_search_provider"] == "tavily"
    assert body["tavily_api_key"] == "***"

    # Masked PUT keeps in-memory key; invalid provider falls back to ddgs
    resp2 = client.put(
        "/settings",
        json={"web_search_provider": "nope", "tavily_api_key": "***"},
    )
    assert resp2.status_code == 200
    assert resp2.json()["web_search_provider"] == "ddgs"
    assert resp2.json()["tavily_api_key"] == "***"


def test_settings_ocr_and_search_persist_across_app_restart(tmp_path, monkeypatch):
    """CLI-style Settings(host, port) must reload OCR / web search from disk."""
    monkeypatch.setenv("LUMINA_DATA_DIR", str(tmp_path))
    monkeypatch.delenv("TAVILY_API_KEY", raising=False)
    monkeypatch.delenv("LUMINA_TAVILY_API_KEY", raising=False)
    monkeypatch.delenv("LUMINA_OCR_CLOUD_API_KEY", raising=False)
    monkeypatch.delenv("LUMINA_OCR_CLOUD_BASE_URL", raising=False)
    monkeypatch.delenv("LUMINA_OCR_CLOUD_MODEL", raising=False)

    app1 = create_app(Settings(host="127.0.0.1", port=17432))
    with TestClient(app1) as client:
        resp = client.put(
            "/settings",
            json={
                "web_search_enabled": False,
                "web_search_provider": "tavily",
                "tavily_api_key": "tvly-restart",
                "ocr_cloud_base_url": "https://example.test/v1",
                "ocr_cloud_model": "vision",
                "ocr_cloud_api_key": "ocr-restart",
                "ocr_cloud_timeout_seconds": 88,
            },
        )
        assert resp.status_code == 200

    app2 = create_app(Settings(host="127.0.0.1", port=17432))
    with TestClient(app2) as client:
        body = client.get("/settings").json()
        assert body["web_search_enabled"] is False
        assert body["web_search_provider"] == "tavily"
        assert body["tavily_api_key"] == "***"
        assert body["ocr_cloud_base_url"] == "https://example.test/v1"
        assert body["ocr_cloud_model"] == "vision"
        assert body["ocr_cloud_api_key"] == "***"
        assert body["ocr_cloud_timeout_seconds"] == 88


def test_settings_secrets_persist_across_reload(client, tmp_path):
    resp = client.put(
        "/settings",
        json={
            "web_search_provider": "tavily",
            "tavily_api_key": "tvly-persist",
            "models": {
                "resources": [
                    {
                        "id": "openai",
                        "provider": "openai",
                        "model": "gpt-4o-mini",
                        "base_url": "https://api.openai.com/v1",
                        "api_key": "sk-persist",
                    }
                ],
                "chat": {"priority": ["openai"]},
                "summarize": {"priority": ["ollama"]},
            },
        },
    )
    assert resp.status_code == 200

    secrets_file = tmp_path / "secrets.json"
    assert secrets_file.exists()
    secrets = json.loads(secrets_file.read_text(encoding="utf-8"))
    assert secrets["resources"]["openai"] == "sk-persist"
    assert secrets["tavily"] == "tvly-persist"

    from lumina_core.settings_store import load_models, load_settings

    reloaded_settings = load_settings(tmp_path)
    reloaded_models = load_models(tmp_path)
    assert reloaded_settings.web_search_provider == "tavily"
    assert reloaded_settings.tavily_api_key == "tvly-persist"
    openai = reloaded_models.resource_by_id("openai")
    assert openai is not None
    assert openai.api_key == "sk-persist"


def test_settings_models_persist_and_redact(client, tmp_path):
    payload = {
        "models": {
            "resources": [
                {
                    "id": "openai",
                    "provider": "openai",
                    "model": "gpt-4o-mini",
                    "base_url": "https://api.openai.com/v1",
                    "api_key": "sk-test-secret",
                    "concurrency": 4,
                },
                {
                    "id": "ollama",
                    "provider": "ollama",
                    "model": "qwen3.5:9b",
                    "base_url": "http://127.0.0.1:11434",
                    "concurrency": 1,
                },
            ],
            "chat": {"priority": ["openai"]},
            "summarize": {"priority": ["ollama"]},
        }
    }
    resp = client.put("/settings", json=payload)
    assert resp.status_code == 200
    body = resp.json()
    openai = next(r for r in body["models"]["resources"] if r["id"] == "openai")
    ollama = next(r for r in body["models"]["resources"] if r["id"] == "ollama")
    assert openai["api_key"] == "***"
    assert openai["model"] == "gpt-4o-mini"
    assert openai["concurrency"] == 4
    assert ollama["model"] == "qwen3.5:9b"
    assert ollama["concurrency"] == 1

    disk = json.loads((tmp_path / "models.json").read_text(encoding="utf-8"))
    disk_openai = next(r for r in disk["resources"] if r["id"] == "openai")
    assert disk_openai["api_key"] is None
    assert "job_concurrency" not in disk
    assert next(r for r in disk["resources"] if r["id"] == "ollama")["model"] == "qwen3.5:9b"

    resp2 = client.put(
        "/settings",
        json={
            "models": {
                **payload["models"],
                "resources": [
                    {**payload["models"]["resources"][0], "api_key": "***", "model": "gpt-4o"},
                    payload["models"]["resources"][1],
                ],
            }
        },
    )
    assert resp2.status_code == 200
    openai2 = next(r for r in resp2.json()["models"]["resources"] if r["id"] == "openai")
    assert openai2["api_key"] == "***"
    assert openai2["model"] == "gpt-4o"
