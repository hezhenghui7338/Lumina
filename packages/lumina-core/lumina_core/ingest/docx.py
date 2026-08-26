"""DOCX ingestion via python-docx."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from lumina_core.chunker.markers import heading_marker


def _paragraph_text(paragraph) -> str:
    return re.sub(r"[ \t]+", " ", paragraph.text).strip()


def load_docx(path: Path) -> tuple[str, dict[str, Any]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx: pip install python-docx") from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ValueError("Invalid or encrypted DOCX document") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = _paragraph_text(paragraph)
        if not text:
            continue
        style_name = (getattr(paragraph.style, "name", "") or "").lower()
        if style_name.startswith("heading") or style_name in {"title", "subtitle"}:
            level = 1
            if style_name.startswith("heading"):
                suffix = style_name.replace("heading", "").strip()
                if suffix.isdigit():
                    level = max(1, int(suffix))
            parts.append(heading_marker(text, level))
        else:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [re.sub(r"\s+", " ", cell.text).strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError("DOCX document contains no readable text")

    properties = document.core_properties
    metadata = {
        key: value.strip()
        for key, value in {
            "title": properties.title,
            "author": properties.author,
        }.items()
        if isinstance(value, str) and value.strip()
    }
    return text, metadata
